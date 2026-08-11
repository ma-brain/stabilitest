#!/usr/bin/env python3
"""Apply the Welch recalibration corrections to the physician's guide."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


UPDATED_SENTINEL = (
    "The prospective Welch study, version welch-2026-2, ended at training "
    "with no_feasible_thresholds."
)


PARAGRAPH_REPLACEMENTS = {
    "That number is the removal fragility index, and it is the heart of the framework. Validation studies (Section 4) show a stark pattern: significant results that are actually chance findings typically collapse after removing only 1-2 patients - about 2% of the sample - while genuine, large treatment effects withstand the removal of 10-14% of patients.":
        "The removal fragility index records how many patients the determined skeptic removed before the conclusion changed. In the historical Welch simulation, null scenarios often changed after 1-2 removals, while large-effect scenarios usually required more. These are group-level patterns, not validated diagnostic cutoffs for an individual trial.",
    "For the simplest and most common setting - comparing two independent groups with the standard (Welch) t-test - this validation exists. Chance findings that happened to reach significance averaged a score of about 52 and typically collapsed after 1-2 adversarial removals. Genuine large effects averaged about 75 and withstood removal of 10-14% of patients. From this separation come the reference ranges: above 70 = Robust; 55 to 70 = Moderately robust; 55 or below = Fragile.":
        "An early Welch simulation suggested that average scores differed between null and large-effect scenarios, but the score distributions overlapped too much to justify clinical reference ranges. The prospective Welch study, version welch-2026-2, ended at training with no_feasible_thresholds. Because no candidate met the pre-specified requirements, the held-out validation data were not opened. The technical archive retains the exact audit identifiers.",
    "A second setting has since been validated too: comparing two groups on a yes/no outcome - responder versus non-responder, say - with Fisher's exact test. Because one yes/no comparison carries less information than a continuous measurement, this validation supports only two tiers, not three: a score of 58 or below is Fragile, above 58 is Not fragile - there is no Robust tier here. The range applies only within the profile it was validated on (roughly 25-200 patients per arm, arms of similar size, and a control-arm event rate between 8% and 55% with at least three events and three non-events) and only under the specific weighting the calibration study used. Outside that profile, or under the package's default weighting, a binary-outcome trial still gets the full numeric breakdown - just no printed verdict.":
        "One setting currently has a validated categorical interpretation: comparing two groups on a yes/no outcome with Fisher's exact test. A score of 58 or below is Fragile and a score above 58 is Not fragile; there is no Robust tier. This mapping applies only within its validated profile (roughly 25-200 patients per arm, arms of similar size, and a control-arm event rate between 8% and 55% with at least three events and three non-events) and only under the specific weighting used by the calibration study. Outside that profile, or under the package's default weighting, the numeric breakdown is still reported but the label is left blank.",
    "Crucially, these ranges were validated only for those two specific analyses. For every other analysis - adjusted comparisons, survival analyses, binary-outcome trials outside the validated profile above, equivalence testing - stabilitest reports all the numbers but shows no verdict label. Like a responsible laboratory, it will not print a reference range it has not validated for your assay. The designers call this “fail-closed”: when validation is missing or fails, the label stays silent.":
        "Welch categorical labels are suppressed; numeric scores and component metrics remain available. The same rule applies to adjusted comparisons, survival analyses, binary-outcome trials outside the validated Fisher profile, and equivalence testing. In plain language, stabilitest still shows the measurements but leaves the categorical label blank when no validated reference range exists. This deliberate fail-closed behaviour avoids presenting an invented diagnosis.",
    "The robustness analysis returns a score of 72.5 - Robust - built from:":
        "The robustness analysis returns a numeric score of 72.5. No categorical Welch verdict is emitted. The score is built from:",
    "Determined skeptic: 6 patients (11% of the sample) had to be removed before significance was lost (p rising step by step from 0.0018 to 0.060). That matches the validated signature of a genuine effect (5-6 at this sample size) and is far from the chance-finding signature of 1-2.":
        "Determined skeptic: 6 patients (11% of the sample) had to be removed before significance was lost (p rising step by step from 0.0018 to 0.060). This is useful descriptive context for reviewing the result and the named patients, but it is not a validated diagnosis of whether the treatment effect is genuine.",
    "What does a clinician do with this? The six named patients - led by one extreme responder whose pain fell by 52 points - deserve source-data review: protocol adherence, concomitant medication, transcription errors. A supplementary rank-based analysis, which is less influenced by extreme values, is also sensible. “Robust” does not mean beyond question; it means the verdict does not hinge on a data quirk - and the analysis has told you exactly which charts to pull to make sure.":
        "What does a clinician do with this? The six named patients - led by one extreme responder whose pain fell by 52 points - deserve source-data review: protocol adherence, concomitant medication, and transcription errors. A supplementary rank-based analysis, which is less influenced by extreme values, is also sensible. The useful result is the component profile and the list of charts to review, not an unsupported Robust label.",
    "6. When validation honestly fails: the adjusted-analysis story":
        "6. When validation honestly fails: Welch and adjusted analyses",
    "Most confirmatory trials do not use a plain two-group comparison; they adjust for each patient's baseline value (ANCOVA). The natural next step was to validate reference ranges for that setting. In 2026, two large pre-registered calibration studies were run - 2,700 simulated significant trials in the first training set alone, with the pass/fail criteria frozen in advance, exactly as one would pre-register a clinical trial.":
        "The prospective Welch study, version welch-2026-2, ended at training with no_feasible_thresholds. It completed 4,500 significant training analyses, but no proposed cutoff was both safe against false reassurance and useful for identifying clear effects. Because no candidate met the pre-specified requirements, the held-out validation data were not opened.",
    "Both attempts reached the same conclusion: no reliable reference ranges exist for the adjusted setting. Follow-up analysis showed why. In a clean adjusted analysis, the robustness score moves in near lockstep with the p-value itself, and distinguishing chance findings from moderate genuine effects at the required error levels would demand a sharpness of discrimination that no score built from the same data can mathematically deliver. It is not that the studies were unlucky; the target itself was unreachable.":
        "The earlier 55/70 Welch bands therefore remain historical exploratory markers, not clinical reference ranges. In prospective training they would have reassured 45.8% of null results and identified only 49.9% of clear-effect results. The package did not loosen its requirements to make the bands pass; it kept the numeric output and left the Welch label blank.",
    "There were two possible responses. One is quietly loosening the passing criteria until something “validates.” The other is saying no. The package says no: adjusted analyses keep their numeric outputs, and the verdict field stays empty. For a physician, this is exactly the behaviour you want from a diagnostic - silence rather than an invented reference range. The negative result is documented in full, with the same rigor as a positive one would have been.":
        "Adjusted analyses tell a similar story. Two pre-registered ANCOVA training studies found no cutoff that met the frozen error requirements, so both stopped before held-out validation. ANCOVA analyses still return numeric scores and stress-test details, but their categorical label is blank. For a physician, that silence is preferable to an invented reference range.",
    "A re-aim was then tried at a target with more direct clinical meaning: instead of a three-zone verdict, could the score detect adjusted analyses whose assumptions had been quietly violated - the one job a p-value genuinely cannot do? A pre-registered study tested exactly this, and again the answer was no: the added detective power over the p-value alone was statistically indistinguishable from zero. That makes three rigorously documented negative findings for the adjusted-analysis setting, each published with the same rigor a positive result would have received, and none used to loosen the original passing criteria. A further re-aim - calibrating the score as an estimated probability that a repeat trial would again be significant, rather than as a Robust/Fragile verdict - is a considered next step that has not yet been attempted.":
        "A later pre-registered ANCOVA study asked whether the score could detect assumption violations better than the p-value alone. The added discrimination was statistically indistinguishable from zero. Across these negative studies, the original requirements were not weakened and numeric outputs were retained. This is what fail-closed means in practice: useful measurements remain visible, while an unsupported categorical diagnosis does not.",
    "How many patients would a determined skeptic need to remove to overturn the finding? If fewer than about 5% of the sample, read the paper carefully - whatever the p-value says.":
        "How many patients would a determined skeptic need to remove to overturn the finding? If fewer than about 5% of the sample, treat that as a prompt for careful clinical and data-quality review, not as a validated diagnostic cutoff.",
    "If a result is labelled Robust or Fragile, were the reference ranges behind that label validated for this exact analysis type - or borrowed from a different one?":
        "If a categorical label is printed, was it validated for this exact analysis type, profile, and weighting? If the label is blank, use the numeric components rather than borrowing a reference range from another method.",
}


TABLE_CELL_REPLACEMENTS = {
    "The smallest number of patients found whose removal makes the result lose significance. One or two = the signature of a chance finding.":
        "The smallest number of patients found whose removal makes the result lose significance. Small values prompt review; they are not a diagnosis of a chance finding.",
    "Validation studies in simulated trials where the truth is known, establishing what scores chance findings versus genuine effects produce.":
        "Studies using simulated trials with known truth to test whether proposed score cutoffs are safe and useful before categorical labels are allowed.",
    "A test comparing two groups on a yes/no outcome. The second setting with a validated verdict - Fragile or Not fragile - within its calibrated profile (Section 4).":
        "A test comparing two groups on a yes/no outcome. It is the only current setting with a validated categorical label - Fragile or Not fragile - within its calibrated profile (Section 4).",
    "The design principle that no verdict label is shown unless its reference ranges were validated for exactly that analysis type.":
        "The design principle that numeric results remain visible but the categorical label is left blank unless cutoffs were validated for that exact analysis type.",
}


PREFIXES = {
    "Determined skeptic: 6 patients (11% of the sample) had to be removed before significance was lost (p rising step by step from 0.0018 to 0.060). That matches the validated signature of a genuine effect (5-6 at this sample size) and is far from the chance-finding signature of 1-2.":
        "Determined skeptic: ",
    "Crucially, these ranges were validated only for those two specific analyses. For every other analysis - adjusted comparisons, survival analyses, binary-outcome trials outside the validated profile above, equivalence testing - stabilitest reports all the numbers but shows no verdict label. Like a responsible laboratory, it will not print a reference range it has not validated for your assay. The designers call this “fail-closed”: when validation is missing or fails, the label stays silent.":
        "Welch categorical labels are suppressed; numeric scores and component metrics remain available. ",
}


POST_UPDATE_REPLACEMENTS = {
    "An early Welch simulation suggested that average scores differed between null and large-effect scenarios, but the score distributions overlapped too much to justify clinical reference ranges. The prospective Welch study, version welch-2026-2, ended at training with no_feasible_thresholds. Its frozen candidate record has hash 9c45481b952cab7cb9b9086e37924a39d83fe0484628745dbe2e79eb33e8797d. Because no candidate met the pre-specified requirements, the held-out validation data were not opened.":
        "An early Welch simulation suggested that average scores differed between null and large-effect scenarios, but the score distributions overlapped too much to justify clinical reference ranges. The prospective Welch study, version welch-2026-2, ended at training with no_feasible_thresholds. Because no candidate met the pre-specified requirements, the held-out validation data were not opened. The technical archive retains the exact audit identifiers.",
    "2.3 Replaying the trial (the bootstrap)":
        "2.3 Bootstrap same-decision rate",
    "Finally, the computer builds thousands of simulated repetitions of the trial by resampling the observed patients, and counts how often the repetition reaches the same conclusion. The result reads like a weather forecast: “92% of simulated reruns of this trial confirm the original verdict.”":
        "Finally, the computer repeatedly resamples the observed patients and counts how often each resample reaches the same significance decision. The result is the bootstrap same-decision rate; for example, 92% means that 92% of these observed-data resamples preserved the original decision.",
    "A caveat here too: this replay percentage mostly restates how far the p-value sits from 0.05. A borderline result gives about 50-60% even in perfectly clean data. It answers “would a repeat trial likely agree?”, which is a different question from “is this result resting on a few patients?” - which is why it is reported separately and weighted lightly.":
        "This rate mostly restates how far the p-value sits from 0.05. A borderline result gives about 50-60% even in perfectly clean data. It is an empirical plug-in quantity under the observed data, not the probability that a new trial will replicate the finding. That is why it is reported separately and weighted lightly.",
    "The three stress tests are combined into a single score from 0 to 100, much as clinical practice combines several measurements into a composite index. By default, the fragility test and the leave-one-out test each contribute 40% and the replay contributes 20%; the weights are stated openly and can be fixed in advance in the trial's statistical analysis plan.":
        "The three stress tests are combined into a single score from 0 to 100, much as clinical practice combines several measurements into a composite index. By default, the fragility test and the leave-one-out test each contribute 40% and the bootstrap same-decision rate contributes 20%; the weights are stated openly and can be fixed in advance in the trial's statistical analysis plan.",
    "Replay: 92% of simulated reruns confirmed the conclusion.":
        "Bootstrap same-decision rate: 92% of observed-data resamples preserved the conclusion.",
}


POST_UPDATE_TABLE_CELL_REPLACEMENTS = {
    "Replay (bootstrap)":
        "Bootstrap same-decision rate",
    "If we could rerun the trial in similar patients, how often would it reach the same conclusion?":
        "How often do resamples from the observed data preserve the original significance decision?",
    "A reproducibility percentage (e.g., 92% of simulated reruns confirm)":
        "A same-decision percentage (for example, 92% of observed-data resamples preserve the decision)",
    "Building thousands of simulated reruns of the trial by resampling the observed patients.":
        "Repeatedly resampling the observed patients and rerunning the analysis.",
    "Reproducibility probability":
        "Bootstrap same-decision rate",
    "The share of those simulated reruns that reach the same conclusion as the original analysis.":
        "The share of observed-data resamples that preserve the original significance decision; not a future-trial replication probability.",
    "A 0-100 summary combining the three stress tests (default emphasis: fragility and leave-one-out 40% each, replay 20%).":
        "A 0-100 summary combining the three stress tests (default emphasis: fragility and leave-one-out 40% each, bootstrap same-decision rate 20%).",
}


def set_paragraph_text(paragraph, new_text: str, prefix: str | None = None) -> None:
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    if prefix and new_text.startswith(prefix) and len(paragraph.runs) > 1:
        paragraph.runs[0].text = prefix
        paragraph.runs[1].text = new_text[len(prefix):]
        for run in paragraph.runs[2:]:
            run.text = ""
        return
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def all_document_text(document) -> str:
    pieces = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            pieces.extend(cell.text for cell in row.cells)
    return "\n".join(pieces)


def update(path: Path) -> None:
    document = Document(path)
    current_text = all_document_text(document)
    changed = False

    paragraph_by_text = {p.text: p for p in document.paragraphs}
    for old, new in POST_UPDATE_REPLACEMENTS.items():
        paragraph = paragraph_by_text.get(old)
        if paragraph is not None:
            set_paragraph_text(paragraph, new)
            print(f"Updated current physician-guide paragraph: {old[:72]}...")
            changed = True

    post_cell_paragraph_by_text = {}
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    post_cell_paragraph_by_text[paragraph.text] = paragraph
    for old, new in POST_UPDATE_TABLE_CELL_REPLACEMENTS.items():
        paragraph = post_cell_paragraph_by_text.get(old)
        if paragraph is not None:
            set_paragraph_text(paragraph, new)
            print(f"Updated current glossary/table cell: {old[:72]}...")
            changed = True

    glossary_heading = next(
        (p for p in document.paragraphs if p.text == "8. Glossary"), None
    )
    if glossary_heading is None:
        raise SystemExit("Expected physician-guide glossary heading was not found")
    if not glossary_heading.paragraph_format.page_break_before:
        glossary_heading.paragraph_format.page_break_before = True
        print("Set the glossary heading to begin on a new page.")
        changed = True

    if UPDATED_SENTINEL in current_text and not any(
        old in current_text
        for old in (*PARAGRAPH_REPLACEMENTS, *TABLE_CELL_REPLACEMENTS)
    ):
        if changed:
            document.save(path)
            print(f"Updated {path}")
        else:
            print("Physician guide is already updated; no changes written.")
        return

    paragraph_by_text = {p.text: p for p in document.paragraphs}
    missing_paragraphs = [
        old for old in PARAGRAPH_REPLACEMENTS if old not in paragraph_by_text
    ]
    if missing_paragraphs:
        raise SystemExit(
            "Expected physician-guide paragraphs were not found:\n- "
            + "\n- ".join(missing_paragraphs)
        )

    cell_paragraph_by_text = {}
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    cell_paragraph_by_text[paragraph.text] = paragraph
    missing_cells = [
        old for old in TABLE_CELL_REPLACEMENTS if old not in cell_paragraph_by_text
    ]
    if missing_cells:
        raise SystemExit(
            "Expected physician-guide table cells were not found:\n- "
            + "\n- ".join(missing_cells)
        )

    for old, new in PARAGRAPH_REPLACEMENTS.items():
        set_paragraph_text(paragraph_by_text[old], new, PREFIXES.get(old))
        print(f"Replaced paragraph: {old[:72]}...")
        changed = True
    for old, new in TABLE_CELL_REPLACEMENTS.items():
        set_paragraph_text(cell_paragraph_by_text[old], new)
        print(f"Replaced glossary cell: {old[:72]}...")
        changed = True

    if changed:
        document.save(path)
        print(f"Updated {path}")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: update-physicians-guide.py PATH.docx")
    path = Path(sys.argv[1])
    if not path.is_file():
        raise SystemExit(f"DOCX not found: {path}")
    update(path)


if __name__ == "__main__":
    main()
